# -*- coding: utf-8 -*-
"""
프로젝트명: KCS 딸깍 기록지
제작자: D
슬로건: 당신의 손목은 소중하니깐..
버전: v6.2 (다운로드 버튼 위치 최적화 - 스크롤 없이 즉시 다운로드 가능하도록 UI 개선)
"""

import io
import os
import re
import openpyxl
from openpyxl.styles import Alignment
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="KCS 딸깍 기록지",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ==============================================================================
# [UI/UX 디자인 시스템] 엔터프라이즈 SaaS 인터페이스 & 모바일 반응형 CSS
# ==============================================================================
st.markdown(
    """
    <style>
    @import url('https://cdn.jsdelivr.net/gh/orioncactus/pretendard/dist/web/static/pretendard.css');

    :root {
        --bg-main: #F8FAFC;
        --card-bg: #FFFFFF;
        --text-primary: #0F172A;
        --text-secondary: #475569;
        --text-muted: #94A3B8;
        --border-color: #E2E8F0;
        --accent-primary: #2563EB;
        --accent-secondary: #4F46E5;
        --shadow-sm: 0 1px 3px rgba(15, 23, 42, 0.04), 0 1px 2px rgba(15, 23, 42, 0.02);
        --shadow-md: 0 10px 25px -5px rgba(15, 23, 42, 0.05), 0 8px 10px -6px rgba(15, 23, 42, 0.02);
    }

    html, body, [class*="css"], div:not([data-testid="stIconMaterial"]) {
        font-family: 'Pretendard', -apple-system, BlinkMacSystemFont, system-ui, Roboto, sans-serif;
        box-sizing: border-box;
    }

    .stApp {
        background-color: var(--bg-main);
        background-image:
            radial-gradient(at 0% 0%, rgba(37, 99, 235, 0.06) 0px, transparent 45%),
            radial-gradient(at 100% 0%, rgba(79, 70, 229, 0.06) 0px, transparent 45%),
            radial-gradient(at 50% 100%, rgba(16, 185, 129, 0.04) 0px, transparent 50%);
        background-attachment: fixed;
    }

    .block-container {
        max-width: 1180px !important;
        padding-top: 2.2rem !important;
        padding-bottom: 4rem !important;
        padding-left: 1.5rem !important;
        padding-right: 1.5rem !important;
    }

    /* 상단 히어로 배너 */
    .hero-card {
        background: linear-gradient(135deg, #0F172A 0%, #1E293B 100%);
        border-radius: 22px;
        padding: 2rem 2.2rem;
        color: #FFFFFF;
        box-shadow: 0 12px 30px -8px rgba(15, 23, 42, 0.22);
        margin-bottom: 1.6rem;
        position: relative;
        overflow: hidden;
        border: 1px solid rgba(255, 255, 255, 0.08);
    }
    .hero-card::after {
        content: "";
        position: absolute;
        top: -50%;
        right: -10%;
        width: 360px;
        height: 360px;
        background: radial-gradient(circle, rgba(59, 130, 246, 0.22) 0%, transparent 70%);
        pointer-events: none;
    }
    .hero-title {
        font-size: clamp(2.1rem, 4.5vw, 3.1rem);
        font-weight: 900;
        line-height: 1.15;
        color: #FFFFFF;
        letter-spacing: -0.03em;
        margin-bottom: 0.55rem;
    }
    .hero-title-accent {
        background: linear-gradient(90deg, #60A5FA 0%, #A7F3D0 100%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
    }
    .hero-desc {
        font-size: clamp(0.88rem, 1.5vw, 0.98rem);
        color: #CBD5E1;
        line-height: 1.6;
        font-weight: 400;
        margin-bottom: 1.25rem;
    }
    .hero-desc b {
        color: #F8FAFC;
        font-weight: 700;
    }

    /* 기기 분류 배지 */
    .badge-wrap {
        display: flex;
        flex-wrap: wrap;
        gap: 8px;
    }
    .chip {
        display: inline-flex;
        align-items: center;
        padding: 0.38rem 0.85rem;
        border-radius: 10px;
        font-size: 0.82rem;
        font-weight: 600;
        backdrop-filter: blur(6px);
        transition: transform 0.15s ease;
    }
    .chip:hover {
        transform: translateY(-1px);
    }
    .chip-code {
        font-weight: 800;
        margin-right: 6px;
        padding: 1px 6px;
        border-radius: 6px;
        font-size: 0.75rem;
    }
    .chip-f { color: #FFEDD5; background: rgba(234, 88, 12, 0.2); border: 1px solid rgba(251, 146, 60, 0.35); }
    .chip-f .chip-code { background: rgba(234, 88, 12, 0.45); color: #FFF7ED; }
    .chip-m { color: #E0E7FF; background: rgba(79, 70, 229, 0.22); border: 1px solid rgba(129, 140, 248, 0.35); }
    .chip-m .chip-code { background: rgba(79, 70, 229, 0.45); color: #EEF2FF; }
    .chip-t { color: #FCE7F3; background: rgba(219, 39, 119, 0.2); border: 1px solid rgba(244, 114, 182, 0.35); }
    .chip-t .chip-code { background: rgba(219, 39, 119, 0.45); color: #FDF2F8; }
    .chip-p { color: #D1FAE5; background: rgba(5, 150, 105, 0.2); border: 1px solid rgba(52, 211, 153, 0.35); }
    .chip-p .chip-code { background: rgba(5, 150, 105, 0.45); color: #ECFDF5; }

    /* 섹션 타이틀 헤더 */
    .section-header {
        display: flex;
        align-items: center;
        justify-content: space-between;
        margin-bottom: 0.75rem;
    }
    .section-title {
        font-size: 1.05rem;
        font-weight: 800;
        color: var(--text-primary);
        display: flex;
        align-items: center;
        gap: 8px;
        letter-spacing: -0.02em;
    }
    .section-pill {
        font-size: 0.76rem;
        font-weight: 700;
        color: var(--accent-primary);
        background: #EFF6FF;
        border: 1px solid #BFDBFE;
        padding: 0.2rem 0.65rem;
        border-radius: 9999px;
    }

    /* Streamlit 기본 Border 컨테이너 모던 카드화 */
    div[data-testid="stVerticalBlockBorderWrapper"] {
        background: var(--card-bg) !important;
        border: 1px solid var(--border-color) !important;
        border-radius: 18px !important;
        padding: 1.35rem 1.4rem !important;
        box-shadow: var(--shadow-md) !important;
    }

    /* 파일 업로더 드롭존 커스텀 */
    div[data-testid="stFileUploader"] {
        width: 100% !important;
        margin-top: 0 !important;
        margin-bottom: 0.85rem !important;
    }
    div[data-testid="stFileUploader"] > label {
        display: none !important;
    }
    div[data-testid="stFileUploader"] section {
        position: relative !important;
        padding: 2.8rem 1.4rem 1.6rem 1.4rem !important;
        border: 2px dashed #94A3B8 !important;
        border-radius: 16px !important;
        background: #F8FAFC !important;
        transition: all 0.2s ease !important;
    }
    div[data-testid="stFileUploader"] section::before {
        content: "📂 교정 성적서(.xlsx, .xlsm)를 이곳에 끌어다 놓거나 선택하세요";
        position: absolute;
        top: 16px;
        left: 12px;
        right: 12px;
        text-align: center;
        font-size: 0.96rem;
        font-weight: 800;
        color: #1E293B;
        letter-spacing: -0.02em;
        pointer-events: none;
    }
    div[data-testid="stFileUploader"] section:hover {
        border-color: var(--accent-primary) !important;
        background: #EFF6FF !important;
    }
    div[data-testid="stFileUploader"] [data-testid="stIconMaterial"] {
        font-family: 'Material Symbols Rounded', 'Material Icons' !important;
    }
    div[data-testid="stFileUploader"] button {
        border-radius: 10px !important;
        font-size: 0.88rem !important;
        font-weight: 700 !important;
    }

    /* 옵션 아이템 행 스타일 */
    .opt-row {
        display: flex;
        align-items: center;
        gap: 12px;
        padding: 0.35rem 0;
    }
    .opt-icon-box {
        width: 40px;
        height: 40px;
        background: #F8FAFC;
        border: 1px solid #E2E8F0;
        border-radius: 11px;
        display: flex;
        align-items: center;
        justify-content: center;
        font-size: 1.15rem;
        flex-shrink: 0;
    }
    .opt-main-lbl {
        font-size: 0.94rem;
        font-weight: 700;
        color: var(--text-primary);
        line-height: 1.25;
    }
    .opt-sub-lbl {
        font-size: 0.77rem;
        color: #64748B;
        margin-top: 2px;
    }
    .opt-divider {
        height: 1px;
        background: #F1F5F9;
        margin: 0.5rem 0;
    }

    /* 양식 선택 라디오 버튼 모던 스타일링 */
    div[data-testid="stRadio"] > div {
        gap: 0.6rem !important;
        margin-top: 0.35rem !important;
    }
    div[data-testid="stRadio"] label {
        background: #F8FAFC;
        border: 1px solid #E2E8F0;
        padding: 0.42rem 0.9rem !important;
        border-radius: 10px;
        font-weight: 700 !important;
        font-size: 0.86rem !important;
        transition: all 0.15s ease;
    }
    div[data-testid="stRadio"] label:hover {
        border-color: var(--accent-primary);
        background: #EFF6FF;
    }

    .info-callout {
        background: #F0F9FF;
        border: 1px solid #BAE6FD;
        border-radius: 12px;
        padding: 0.8rem 1rem;
        margin-top: 0.9rem;
        font-size: 0.8rem;
        color: #0369A1;
        line-height: 1.5;
        display: flex;
        gap: 8px;
        align-items: flex-start;
    }

    /* 액션 버튼 스타일링 */
    div.action-btn-del button {
        height: 50px !important;
        border: 1px solid #CBD5E1 !important;
        background: #FFFFFF !important;
        color: #475569 !important;
        font-size: 0.94rem !important;
        font-weight: 700 !important;
        border-radius: 12px !important;
        transition: all 0.2s ease !important;
    }
    div.action-btn-del button:hover {
        color: #DC2626 !important;
        background: #FEF2F2 !important;
        border-color: #FCA5A5 !important;
    }

    div.action-btn-run button {
        height: 50px !important;
        font-size: 0.95rem !important;
        font-weight: 700 !important;
        border-radius: 12px !important;
        transition: all 0.2s ease !important;
    }
    div.action-btn-run button:disabled {
        background-color: #F8FAFC !important;
        color: #94A3B8 !important;
        border: 1px solid #E2E8F0 !important;
        cursor: not-allowed !important;
        opacity: 0.85 !important;
        box-shadow: none !important;
    }

    /* 메트릭 카드 디자인 */
    div[data-testid="stMetric"] {
        background: #FFFFFF !important;
        border: 1px solid #E2E8F0 !important;
        border-radius: 16px !important;
        padding: 1rem 1.15rem !important;
        box-shadow: var(--shadow-sm) !important;
        min-width: 0 !important;
        width: 100% !important;
    }
    div[data-testid="stMetricLabel"] {
        font-size: 0.82rem !important;
        font-weight: 700 !important;
        color: #64748B !important;
    }
    div[data-testid="stMetricValue"] {
        font-size: 1.38rem !important;
        font-weight: 800 !important;
        color: #0F172A !important;
        white-space: normal !important;
        word-break: break-all !important;
        line-height: 1.25 !important;
        margin-top: 0.25rem !important;
    }

    /* 다운로드 버튼 강조 */
    div[data-testid="stDownloadButton"] button {
        height: 54px !important;
        border-radius: 14px !important;
        font-size: 1.05rem !important;
        font-weight: 800 !important;
        background: linear-gradient(135deg, #2563EB 0%, #4F46E5 100%) !important;
        color: #FFFFFF !important;
        border: none !important;
        box-shadow: 0 8px 20px -4px rgba(37, 99, 235, 0.38) !important;
        transition: all 0.2s ease !important;
    }
    div[data-testid="stDownloadButton"] button:hover {
        transform: translateY(-1px) !important;
        box-shadow: 0 12px 24px -4px rgba(37, 99, 235, 0.48) !important;
    }

    /* 모바일 및 태블릿 반응형 최적화 */
    @media (max-width: 768px) {
        .block-container {
            padding-top: 1rem !important;
            padding-left: 0.85rem !important;
            padding-right: 0.85rem !important;
            padding-bottom: 2.5rem !important;
        }
        .hero-card {
            padding: 1.35rem 1.2rem;
            border-radius: 16px;
            margin-bottom: 1.1rem;
        }
        .hero-title {
            font-size: 1.75rem;
        }
        .hero-desc {
            font-size: 0.84rem;
            margin-bottom: 1rem;
        }
        .chip {
            padding: 0.3rem 0.65rem;
            font-size: 0.75rem;
        }
        div[data-testid="stFileUploader"] section {
            padding: 2.6rem 0.9rem 1.2rem 0.9rem !important;
        }
        div[data-testid="stFileUploader"] section::before {
            font-size: 0.84rem;
            top: 12px;
        }
        div[data-testid="stMetric"] {
            padding: 0.8rem 0.9rem !important;
        }
        div[data-testid="stMetricValue"] {
            font-size: 1.15rem !important;
        }
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# ==============================================================================
# [백엔드 유틸리티 및 정밀 포맷팅 함수]
# ==============================================================================

if "uploader_key" not in st.session_state:
    st.session_state["uploader_key"] = 0

# 신규 기록지 양식(1슬롯 = 6행, 가운데 3~4번째 행 병합)에 맞춘 포인트별 행 오프셋
# 1~5번째 측정점: 첫 번째 슬롯의 +0, +1, +2(병합상단), +4, +5 행
# 6~10번째 측정점: 두 번째 슬롯의 +6, +7, +8(병합상단), +10, +11 행
PT_ROW_OFFSETS = [0, 1, 2, 4, 5, 6, 7, 8, 10, 11]


def clean_text(v):
    """보이지 않는 공백 문자(한글 채움 문자, Zero-width space 등) 및 수식 잔여물 제거"""
    if v is None:
        return ""
    s = str(v)
    s = re.sub(r"[\u3164\u200b\u200c\u200d\ufeff\xa0]", " ", s).strip()
    if s.startswith("="):
        return ""
    return s


def get_decimal_places(val):
    """숫자 또는 문자열에서 유효 소수점 자릿수 계산 (부동소수점 오차 보정 포함)"""
    if val is None or str(val).strip() in ("", "-"):
        return 0
    try:
        s = str(val).strip()
        if "." in s:
            dec = s.split(".")[1]
            if len(dec) > 5 and ("9999" in dec or "0000" in dec):
                s = f"{round(float(s), 4)}"
                if "." in s:
                    return len(s.split(".")[1].rstrip("0"))
                return 0
            return len(dec)
    except Exception:
        pass
    return 0


def format_number(val, decimals):
    """지정된 소수점 자릿수로 반올림 포맷팅"""
    if val is None or val == "" or val == "-":
        return "-"
    try:
        f = float(val)
        if decimals == 0:
            return f"{int(round(f))}"
        return f"{f:.{decimals}f}"
    except Exception:
        return str(val).replace(",", "").strip()


def format_meaningful_decimals(val):
    """무의미한 소수점(.00)을 제거하고 유의미한 값만 문자열로 반환"""
    if val is None or str(val).strip() in ("", "-"):
        return "-"
    try:
        f_val = float(str(val).replace(",", "").strip())
        if f_val.is_integer():
            return str(int(f_val))
        return f"{f_val:.6f}".rstrip("0").rstrip(".")
    except Exception:
        return clean_text(val)


def get_formatted_value(cell):
    """Excel 셀 서식(number_format)을 분석하여 화면에 표시되는 소수점 그대로 문자열 반환"""
    if cell is None or cell.value is None or str(cell.value).strip() in ("", "-"):
        return None
    val = cell.value
    val_str = clean_text(val)

    if isinstance(val, (int, float)):
        fmt = str(cell.number_format or "General")
        if "." in fmt:
            dec_part = fmt.split(".")[1].split(";")[0]
            decimals = sum(1 for ch in dec_part if ch in "0#?")
            val_str = f"{val:.{decimals}f}"
        elif fmt != "General" and float(val).is_integer():
            val_str = f"{int(val)}"
        else:
            val_str = format_meaningful_decimals(val)

    return val_str


def format_sub_weight(val):
    """호퍼 저울 가분동 정수 포맷팅"""
    if val is None or val == "" or val == "-":
        return "-"
    try:
        return f"{int(round(float(val)))}"
    except Exception:
        return str(val).split(".")[0].replace(",", "").strip()


def extract_set_count(text):
    """텍스트에서 '(X조)' 또는 'X조' 숫자 추출"""
    if not text:
        return None
    m = re.search(r"(\d+)\s*조", str(text))
    return int(m.group(1)) if m else None


def clean_sensor_type(val):
    """온도계 센서 타입에서 불필요한 None, 잉여 괄호, 공백 등을 깔끔하게 정제"""
    if not val:
        return ""
    if "(단위" in val:
        val = val.split("(단위")[0]
    val = val.replace("None", "").strip()
    while val.count(")") > val.count("("):
        val = val.rsplit(")", 1)[0].strip()
    val = re.sub(r"\s+[tT]ype", "", val).strip()
    val = re.sub(r"\(\s+", "(", val)
    val = re.sub(r"\s+\)", ")", val)
    return val


def format_capacity_and_resolution_robust(cap_raw, res_raw):
    """저울 최대용량(MAX) 및 최소눈금(DIV) 단일/듀얼 레인지 텍스트 표준화"""
    c_str = clean_text(cap_raw)
    r_str = clean_text(res_raw)

    c_norm = re.sub(r"(\d+)\s+(\d{3})(?!\d)", r"\1\2", c_str)
    r_norm = re.sub(r"(\d+)\s+(\d{3})(?!\d)", r"\1\2", r_str)
    c_norm_check = re.sub(r"(\d+),(\d{3})(?!\d)", r"\1\2", c_norm)

    is_dual = False
    if ("," in r_norm or "/" in r_norm) and len(re.findall(r"[\d\.]+", r_norm)) >= 2:
        is_dual = True
    elif "," in c_norm and len(re.findall(r"[\d\.]+", c_norm_check)) >= 2:
        is_dual = True

    if is_dual:
        cap_nums = re.findall(r"([\d\.]+)", c_norm)
        cap_units = re.findall(r"(kg|g|mg|t)", c_norm)
        unit_c = cap_units[-1] if cap_units else "g"
        if len(cap_nums) >= 2:
            c1 = f"{int(float(cap_nums[0])) if float(cap_nums[0]).is_integer() else cap_nums[0]}"
            c2 = f"{int(float(cap_nums[1])) if float(cap_nums[1]).is_integer() else cap_nums[1]}"
            max_str = f"{c1} / {c2} {unit_c}"
        else:
            max_str = c_norm

        res_nums = re.findall(r"([\d\.]+)", r_norm)
        res_units = re.findall(r"(kg|g|mg|t)", r_norm)
        unit_r = res_units[-1] if res_units else "mg"
        if len(res_nums) >= 2:
            r1 = f"{int(float(res_nums[0])) if float(res_nums[0]).is_integer() else res_nums[0]}"
            r2 = f"{int(float(res_nums[1])) if float(res_nums[1]).is_integer() else res_nums[1]}"
            div_str = f"{r1} / {r2} {unit_r}"
        else:
            div_str = r_norm
    else:
        cap_nums = re.findall(r"([\d\.]+)", c_norm_check)
        cap_units = re.findall(r"(kg|g|mg|t)", c_norm_check)
        unit_c = cap_units[-1] if cap_units else "kg"
        if cap_nums:
            f = float(cap_nums[0])
            max_str = f"{int(f) if f.is_integer() else f} {unit_c}"
        else:
            max_str = c_norm_check

        res_nums = re.findall(r"([\d\.]+)", r_norm)
        res_units = re.findall(r"(kg|g|mg|t)", r_norm)
        unit_r = res_units[-1] if res_units else "kg"
        if res_nums:
            f = float(res_nums[0])
            div_str = f"{int(f) if f.is_integer() else f} {unit_r}"
        else:
            div_str = r_norm

    return max_str, div_str, is_dual


def safe_load_workbook(uploaded_file, data_only=True):
    """업로드된 엑셀 스트림을 안전하게 로드"""
    uploaded_file.seek(0)
    try:
        return openpyxl.load_workbook(uploaded_file, data_only=data_only)
    except Exception:
        uploaded_file.seek(0)
        return openpyxl.load_workbook(uploaded_file, data_only=data_only, read_only=True)


# ==============================================================================
# [상단 히어로 헤더 & 메인 레이아웃]
# ==============================================================================
st.markdown(
    """
    <div class="hero-card">
        <div class="hero-title">
            KCS <span class="hero-title-accent">딸깍 기록지</span>
        </div>
        <div class="hero-desc">
            교정 성적서 엑셀 파일을 업로드하면 즉시 자동 변환합니다.<br>
            ⚠️ 변환 결과에 일부 차이가 있을 수 있으니 최종 확인을 권장드리며, 오류나 업데이트 문의는 <b>D</b>에게 전달해 주세요.
        </div>
        <div class="badge-wrap">
            <div class="chip chip-f"><span class="chip-code">FXX</span>힘 시험기</div>
            <div class="chip chip-m"><span class="chip-code">MXX</span>전기식·호퍼 저울</div>
            <div class="chip chip-t"><span class="chip-code">TXX</span>온도 지시·기록계</div>
            <div class="chip chip-p"><span class="chip-code">PXX</span>압력계·차압계</div>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)

col_left, col_right = st.columns([1.45, 1.0], gap="large")

with col_right:
    with st.container(border=True):
        st.markdown(
            """
            <div class="section-header">
                <div class="section-title">⚙️ 변환 및 출력 옵션</div>
                <span class="section-pill">맞춤 설정</span>
            </div>
            """,
            unsafe_allow_html=True,
        )

        # 기록지 양식 선택 (10칸 / 12칸)
        c_ico_t, c_lbl_t = st.columns([0.7, 4.4])
        with c_ico_t:
            st.markdown('<div class="opt-icon-box" style="color: #2563EB;">📋</div>', unsafe_allow_html=True)
        with c_lbl_t:
            st.markdown(
                '<div class="opt-main-lbl">기록지 양식 선택</div><div class="opt-sub-lbl">페이지당 출력할 기록지 칸 수 (10칸 / 12칸)</div>',
                unsafe_allow_html=True,
            )
            slot_type = st.radio(
                "기록지 양식 선택",
                options=["10칸 기록지", "12칸 기록지"],
                index=0,
                horizontal=True,
                label_visibility="collapsed",
            )

        MAX_SLOTS_PER_PAGE = 12 if "12칸" in slot_type else 10
        TEMPLATE_FILE = "12칸 기록지.xlsx" if MAX_SLOTS_PER_PAGE == 12 else "10칸 기록지.xlsx"

        st.markdown('<div class="opt-divider"></div>', unsafe_allow_html=True)

        c_ico0, c_lbl0, c_sw0 = st.columns([0.7, 3.3, 1.1])
        with c_ico0:
            st.markdown('<div class="opt-icon-box" style="color: #4F46E5;">⚡</div>', unsafe_allow_html=True)
        with c_lbl0:
            st.markdown(
                '<div class="opt-main-lbl">자동 작업 모드</div><div class="opt-sub-lbl">파일 업로드 시 즉시 자동 변환</div>',
                unsafe_allow_html=True,
            )
        with c_sw0:
            auto_mode = st.toggle("auto_mode", value=True, label_visibility="collapsed")

        st.markdown('<div class="opt-divider"></div>', unsafe_allow_html=True)

        c_ico1, c_lbl1, c_sw1 = st.columns([0.7, 3.3, 1.1])
        with c_ico1:
            st.markdown('<div class="opt-icon-box" style="color: #2563EB;">#</div>', unsafe_allow_html=True)
        with c_lbl1:
            st.markdown(
                '<div class="opt-main-lbl">넘버(No) 비우기</div><div class="opt-sub-lbl">기록지 A열 일련번호 공란 출력</div>',
                unsafe_allow_html=True,
            )
        with c_sw1:
            del_number = st.toggle("del_num", value=False, label_visibility="collapsed")

        st.markdown('<div class="opt-divider"></div>', unsafe_allow_html=True)

        c_ico2, c_lbl2, c_sw2 = st.columns([0.7, 3.3, 1.1])
        with c_ico2:
            st.markdown('<div class="opt-icon-box" style="color: #0D9488;">📐</div>', unsafe_allow_html=True)
        with c_lbl2:
            st.markdown(
                '<div class="opt-main-lbl">명목값(범위) 비우기</div><div class="opt-sub-lbl">기록지 G열 측정범위 공란 출력</div>',
                unsafe_allow_html=True,
            )
        with c_sw2:
            del_nominal = st.toggle("del_nom", value=False, label_visibility="collapsed")

        st.markdown('<div class="opt-divider"></div>', unsafe_allow_html=True)

        c_ico3, c_lbl3, c_sw3 = st.columns([0.7, 3.3, 1.1])
        with c_ico3:
            st.markdown('<div class="opt-icon-box" style="color: #10B981;">📊</div>', unsafe_allow_html=True)
        with c_lbl3:
            st.markdown(
                '<div class="opt-main-lbl">측정데이터 비우기</div><div class="opt-sub-lbl">기록지 H열 측정결과 공란 출력</div>',
                unsafe_allow_html=True,
            )
        with c_sw3:
            del_data = st.toggle("del_dat", value=False, label_visibility="collapsed")

        st.markdown(
            """
            <div class="info-callout">
                <span>💡</span>
                <span>특정 항목을 현장에서 수기로 직접 기입해야 할 때 해당 비우기 옵션을 켜면 빈칸으로 깔끔하게 생성됩니다.</span>
            </div>
            """,
            unsafe_allow_html=True,
        )

with col_left:
    with st.container(border=True):
        uploaded_files = st.file_uploader(
            "교정 성적서를 끌어다 놓으세요",
            type=["xlsx", "xlsm"],
            accept_multiple_files=True,
            key=f"file_uploader_{st.session_state['uploader_key']}",
            help=".xlsx, .xlsm 형식의 교정 성적서 다중 업로드를 지원합니다.",
        )

        file_cnt = len(uploaded_files) if uploaded_files else 0
        st.markdown(
            f"""
            <div class="section-header" style="margin-top: 0.2rem; margin-bottom: 0.9rem;">
                <div class="section-title">📥 업로드 대기열</div>
                <span class="section-pill">총 {file_cnt}개 파일 선택됨</span>
            </div>
            """,
            unsafe_allow_html=True,
        )

        c_del, c_run = st.columns([1, 1.2], gap="small")
        with c_del:
            st.markdown('<div class="action-btn-del">', unsafe_allow_html=True)
            if st.button("🗑️ 전체 초기화", use_container_width=True, help="업로드된 모든 파일을 즉시 비웁니다."):
                st.session_state["uploader_key"] += 1
                st.rerun()
            st.markdown("</div>", unsafe_allow_html=True)

        with c_run:
            st.markdown('<div class="action-btn-run">', unsafe_allow_html=True)
            if auto_mode:
                manual_run = st.button("⚡ 자동 변환 활성화됨", use_container_width=True, disabled=True)
            else:
                manual_run = st.button("⚡ 기록지 변환 시작", use_container_width=True, type="primary")
            st.markdown("</div>", unsafe_allow_html=True)
            
        # 다운로드 버튼이 렌더링될 자리 (스크롤 없이 바로 클릭 가능한 위치)
        download_placeholder = st.empty()

bottom_result_container = st.container()

# ==============================================================================
# [실행부: 엑셀 파싱 및 기록지 생성 로직]
# ==============================================================================

should_run = (auto_mode and uploaded_files) or (not auto_mode and uploaded_files and manual_run)

if not os.path.exists(TEMPLATE_FILE):
    st.error(f"⚠️ 작업 폴더에 '{TEMPLATE_FILE}' 양식 파일이 존재하지 않습니다. 파일을 확인해 주세요.")
elif should_run:
    with st.spinner(f"⚡ 성적서 데이터를 정밀 분석하여 {MAX_SLOTS_PER_PAGE}칸 기록지를 생성하고 있습니다..."):
        try:
            all_instruments = []
            global_client_name = ""
            error_files = []

            for file_idx, file in enumerate(uploaded_files, start=1):
                fname = os.path.basename(file.name).upper()
                wb_src = None
                try:
                    wb_src = safe_load_workbook(file, data_only=True)

                    raw_device_name = ""
                    client_name = ""
                    cert_no = ""
                    maker_name = ""
                    model_name = ""
                    serial_no = ""
                    total_sets = None
                    capacity_raw = ""
                    scale_raw = ""

                    # 1. 기본정보 시트 탐색
                    info_sheet_names = [s for s in wb_src.sheetnames if "기본정보" in s and "영문" not in s]
                    if info_sheet_names:
                        ws_info = wb_src[info_sheet_names[0]]
                        for r_idx, row in enumerate(ws_info.iter_rows(values_only=True)):
                            if r_idx > 15:
                                break
                            for val in row[:6]:
                                cell_v = str(val or "")
                                if "<Calibration_Certificate>" in cell_v:
                                    m_c = re.search(r'<ClientInfo Name=[\'"]([^\'"]+)', cell_v)
                                    if m_c:
                                        client_name = clean_text(m_c.group(1))
                                    m_dev = re.search(r'<InstrumentInfo[^>]*Description=[\'"]([^\'"]+)', cell_v)
                                    if m_dev:
                                        raw_device_name = clean_text(m_dev.group(1))
                                    m_m = re.search(r'<InstrumentInfo[^>]*Maker=[\'"]([^\'"]+)', cell_v)
                                    if m_m:
                                        maker_name = clean_text(m_m.group(1))
                                    m_mod = re.search(r'ManufacturerModel=[\'"]([^\'"]+)', cell_v)
                                    if m_mod and "/" in m_mod.group(1):
                                        model_name = clean_text(m_mod.group(1).split("/")[1])
                                    m_s = re.search(r'SerialNumber=[\'"]([^\'"]+)', cell_v)
                                    if m_s:
                                        serial_no = clean_text(m_s.group(1))
                                    m_cap = re.search(r'Capacity=[\'"]([^\'"]+)', cell_v)
                                    if m_cap:
                                        capacity_raw = clean_text(m_cap.group(1))
                                        total_sets = extract_set_count(capacity_raw)
                                    m_scale = re.search(r'Scale=[\'"]([^\'"]+)', cell_v)
                                    if m_scale:
                                        scale_raw = clean_text(m_scale.group(1))

                        for row in ws_info.iter_rows(values_only=True):
                            row_str = " ".join([str(c or "") for c in row[:35]])
                            if not client_name and "기 관 명" in row_str:
                                client_name = clean_text(row[6] if len(row) > 6 else "")
                            if not cert_no and "성적서번호" in row_str:
                                cert_no = clean_text(row[6] if len(row) > 6 else "")
                            if not maker_name and "제작회사" in row_str:
                                maker_name = clean_text(row[6] if len(row) > 6 else "")
                            if not raw_device_name and "기 기 명" in row_str:
                                raw_device_name = clean_text(row[6] if len(row) > 6 else "")
                            if not model_name and "형식·모델" in row_str:
                                val_m = clean_text(
                                    row[26] if len(row) > 26 and row[26] else (row[6] if len(row) > 6 else "")
                                )
                                if val_m and val_m != "#VALUE!":
                                    model_name = val_m
                            if not capacity_raw and "측정범위" in row_str:
                                capacity_raw = clean_text(row[6] if len(row) > 6 else "")
                            if total_sets is None and ("측정범위" in row_str or "형식·모델" in row_str):
                                total_sets = extract_set_count(row[6] if len(row) > 6 else "")

                    # 성적서 번호 폴백 탐색
                    if not cert_no:
                        for sname in wb_src.sheetnames:
                            ws_temp = wb_src[sname]
                            for r in range(1, 15):
                                for c in range(1, min(25, ws_temp.max_column + 1)):
                                    v_c = str(ws_temp.cell(r, c).value or "")
                                    m_cert = re.search(r"([A-Za-z]\d{2}-\d+-\d+)", v_c)
                                    if m_cert:
                                        cert_no = m_cert.group(1)
                                        break
                                if cert_no:
                                    break
                            if cert_no:
                                break

                    if client_name and not global_client_name:
                        global_client_name = client_name

                    file_instruments = []

                    # ========================================================
                    # [기기 판별 A] 자동 호퍼 저울
                    # ========================================================
                    if "호퍼" in raw_device_name:
                        ws_rec = wb_src["기록지"]
                        all_rec_rows = list(ws_rec.iter_rows(values_only=True))

                        if not total_sets:
                            max_num = 0
                            for r_idx, r_vals in enumerate(all_rec_rows):
                                v60 = str(r_vals[59] if len(r_vals) > 59 else "").strip()
                                if v60.startswith("(") and v60.endswith(")") and v60[1:-1].isdigit():
                                    if (
                                        r_idx + 1 < len(all_rec_rows)
                                        and len(all_rec_rows[r_idx + 1]) > 6
                                        and all_rec_rows[r_idx + 1][6] is not None
                                    ):
                                        max_num = max(max_num, int(v60[1:-1]))
                            total_sets = max_num if max_num > 0 else 3

                        clean_dev = raw_device_name.replace("자동호퍼저울", "자동 호퍼 저울").strip()
                        clean_dev = re.sub(r"\(\d+\s*조\)", "", clean_dev).strip()
                        dev_disp = f"{clean_dev}\n<{total_sets}조>"

                        for r_idx, r_vals in enumerate(all_rec_rows):
                            val_60 = str(r_vals[59] if len(r_vals) > 59 else "").strip()
                            if val_60.startswith("(") and val_60.endswith(")") and val_60[1:-1].isdigit():
                                num = int(val_60[1:-1])
                                if num > total_sets:
                                    continue
                                r_info_idx = r_idx + 1
                                if r_info_idx >= len(all_rec_rows):
                                    continue

                                r_info = all_rec_rows[r_info_idx]
                                r_info_next = (
                                    all_rec_rows[r_info_idx + 1] if r_info_idx + 1 < len(all_rec_rows) else []
                                )

                                cap = r_info[6] if len(r_info) > 6 else 0
                                gubun = clean_text(r_info[16] if len(r_info) > 16 else "")
                                res = r_info_next[6] if len(r_info_next) > 6 else 0
                                serial = clean_text(r_info_next[16] if len(r_info_next) > 16 else serial_no)

                                pts = []
                                for offset in range(2, 16):
                                    curr_idx = r_info_idx + offset
                                    if curr_idx < len(all_rec_rows):
                                        curr_row = all_rec_rows[curr_idx]
                                        nom = curr_row[32] if len(curr_row) > 32 else None
                                        sub = curr_row[36] if len(curr_row) > 36 else None
                                        use_cnt = str(curr_row[37] if len(curr_row) > 37 else "")
                                        ind = curr_row[49] if len(curr_row) > 49 else None

                                        v60_check = str(curr_row[59] if len(curr_row) > 59 else "").strip()
                                        if (
                                            v60_check.startswith("(")
                                            and v60_check.endswith(")")
                                            and v60_check[1:-1].isdigit()
                                        ):
                                            break

                                        if nom is not None and isinstance(nom, (int, float)):
                                            if float(nom) == 0:
                                                continue
                                            has_sub = (
                                                (sub is not None)
                                                and (str(sub).strip() != "")
                                                and ("회 사용" not in use_cnt)
                                            )
                                            pts.append(
                                                {
                                                    "nom": nom,
                                                    "ind": ind if ind is not None else "-",
                                                    "sub": sub if has_sub else "-",
                                                }
                                            )
                                            if len(pts) >= 10:
                                                break

                                file_instruments.append(
                                    {
                                        "type_kind": "호퍼 저울",
                                        "cert_no": cert_no,
                                        "device_name": dev_disp,
                                        "maker": maker_name if maker_name else "대성계기",
                                        "model": "",
                                        "serial": serial,
                                        "max_str": f"{int(float(cap))} kg"
                                        if float(cap).is_integer()
                                        else f"{float(cap):.1f} kg",
                                        "div_str": format_number(res, get_decimal_places(res)) + " kg",
                                        "res_dec": get_decimal_places(res),
                                        "gubun": gubun,
                                        "pts": pts,
                                        "is_sub_channel": num > 1,
                                    }
                                )

                    # ========================================================
                    # [기기 판별 B] 힘 시험기
                    # ========================================================
                    elif fname.startswith("F") or any(
                        k in raw_device_name for k in ["시험기", "인장", "압축", "만능"]
                    ):
                        result_sheets = [
                            s for s in wb_src.sheetnames if s.startswith("교정결과") and "영문" not in s
                        ]

                        if result_sheets:
                            ws_res0 = wb_src[result_sheets[0]]
                            for r_idx, row in enumerate(ws_res0.iter_rows(values_only=True)):
                                if r_idx > 15:
                                    break
                                row_str = " ".join([str(c or "") for c in row if c is not None])
                                for c in row:
                                    if c and re.search(r"[A-Za-z]\d{2}-\d+-\d+", str(c)):
                                        cert_no = clean_text(c)
                                if ("기기명" in row_str or "기      기" in row_str) and not raw_device_name:
                                    for c in row:
                                        if c and any(k in str(c) for k in ["시험기", "인장", "압축"]):
                                            raw_device_name = clean_text(c)
                                if "제작회사 및 형식" in row_str or "제작회사" in row_str:
                                    for c in row:
                                        if c and "/" in str(c):
                                            parts = str(c).split("/")
                                            maker_name = clean_text(parts[0])
                                            model_name = ""
                                if ("기기번호" in row_str or "기   기   번" in row_str) and not serial_no:
                                    for c in row:
                                        if (
                                            c
                                            and any(ch.isdigit() for ch in str(c))
                                            and "기기" not in str(c)
                                            and "*" not in str(c)
                                        ):
                                            serial_no = clean_text(c)

                        if not raw_device_name:
                            raw_device_name = "인장 및 압축 시험기"

                        rec_sheet_names = [s for s in wb_src.sheetnames if "기록지" in s and "영문" not in s]

                        def _sort_rec_key(name):
                            m = re.search(r"\((\d+)\)", name)
                            return int(m.group(1)) if m else 0

                        rec_sheet_names.sort(key=_sort_rec_key)

                        rec_stage_params = []
                        for rs_name in rec_sheet_names:
                            ws_rec_stage = wb_src[rs_name]
                            cap_val = None
                            cap_unit = "kN"
                            div_val = None
                            div_unit = "kN"

                            for r_i, r_vals in enumerate(ws_rec_stage.iter_rows(values_only=True), 1):
                                if r_i > 30:
                                    break
                                row_t = " ".join([str(c or "") for c in r_vals if c is not None])
                                if "최대용량" in row_t:
                                    for c_i, v in enumerate(r_vals):
                                        if v is not None and isinstance(v, (int, float)) and v > 0:
                                            cap_val = v
                                            for u_c in r_vals[c_i + 1 : c_i + 8]:
                                                if u_c and any(
                                                    u in str(u_c) for u in ["kN", "N", "tf", "kgf"]
                                                ):
                                                    cap_unit = (
                                                        str(u_c).replace("(", "").replace(")", "").strip()
                                                    )
                                                    break
                                            break
                                if "분해능" in row_t or "분 해 능" in row_t:
                                    for c_i, v in enumerate(r_vals):
                                        if v is not None and isinstance(v, (int, float)) and v > 0:
                                            div_val = v
                                            for u_c in r_vals[c_i + 1 : c_i + 8]:
                                                if u_c and any(
                                                    u in str(u_c) for u in ["kN", "N", "tf", "kgf"]
                                                ):
                                                    div_unit = (
                                                        str(u_c).replace("(", "").replace(")", "").strip()
                                                    )
                                                    break
                                            break

                            cap_str_rec = (
                                f"{int(cap_val) if float(cap_val).is_integer() else cap_val} {cap_unit}"
                                if cap_val is not None
                                else ""
                            )
                            div_str_rec = f"{div_val} {div_unit}" if div_val is not None else ""
                            rec_stage_params.append({"max_cap": cap_str_rec, "div_str": div_str_rec})

                        stages = []
                        for sname in result_sheets:
                            ws = wb_src[sname]
                            rows = list(ws.iter_rows(values_only=True))
                            r = 0
                            while r < len(rows):
                                row_str = " ".join([str(c or "") for c in rows[r]])
                                if any(
                                    m in row_str
                                    for m in ["압 축 교 정", "압축교정", "인 장 교 정", "인장교정"]
                                ):
                                    method = "인장" if ("인장" in row_str or "인 장" in row_str) else "압축"
                                    start_data_r = -1
                                    unit_str = "kN"

                                    for dr in range(r + 1, min(r + 6, len(rows))):
                                        dr_str = " ".join([str(c or "") for c in rows[dr]])
                                        if "지시하중" in dr_str:
                                            if dr + 1 < len(rows):
                                                u_str = " ".join([str(c or "") for c in rows[dr + 1]])
                                                m_u = re.search(r"\(([A-Za-z]+)\)", u_str)
                                                if m_u:
                                                    unit_str = m_u.group(1)
                                            start_data_r = dr + 2
                                            break

                                    if start_data_r != -1:
                                        pts = []
                                        curr_r = start_data_r
                                        res_val_str = ""
                                        while curr_r < len(rows):
                                            crow = rows[curr_r]
                                            crow_str = " ".join([str(c or "") for c in crow])
                                            if any(
                                                k in crow_str
                                                for k in [
                                                    "상대영점",
                                                    "분해능",
                                                    "상대측정불확도",
                                                    "압 축 교 정",
                                                    "인 장 교 정",
                                                    "KCS-QP",
                                                ]
                                            ):
                                                break
                                            nom = crow[3] if len(crow) > 3 else None
                                            err = crow[28] if len(crow) > 28 else None
                                            if nom is not None:
                                                try:
                                                    f_nom = float(str(nom).replace(",", "").strip())
                                                    if f_nom > 0:
                                                        err_str = "-"
                                                        if err is not None and str(err).strip() != "-":
                                                            f_err = float(err)
                                                            err_str = f"{f_err:+.2f} %"
                                                        pts.append(
                                                            {"nom": f_nom, "ind": err_str, "sub": "-"}
                                                        )
                                                except Exception:
                                                    pass
                                            curr_r += 1

                                        for sr in range(curr_r, min(curr_r + 6, len(rows))):
                                            sr_str = " ".join([str(c or "") for c in rows[sr]])
                                            m_res = re.search(
                                                r"분해능\s*:\s*([\d\.]+\s*[A-Za-z]*)", sr_str
                                            )
                                            if m_res:
                                                res_val_str = m_res.group(1).strip()
                                                if not any(ch.isalpha() for ch in res_val_str):
                                                    res_val_str = f"{res_val_str} {unit_str}"
                                                break

                                        max_nom = max([p["nom"] for p in pts]) if pts else 0
                                        stages.append(
                                            {
                                                "method": method,
                                                "unit": unit_str,
                                                "fallback_max": f"{int(max_nom) if float(max_nom).is_integer() else max_nom} {unit_str}",
                                                "fallback_div": res_val_str
                                                if res_val_str
                                                else f"0.01 {unit_str}",
                                                "pts": pts,
                                            }
                                        )
                                        r = curr_r
                                        continue
                                r += 1

                        total_stages = len(stages)
                        for s_idx, stg in enumerate(stages):
                            if total_stages > 1:
                                dev_disp = (
                                    f"{raw_device_name}\n<{s_idx + 1}단/{total_stages}단>"
                                    if s_idx == 0
                                    else ""
                                )
                            else:
                                dev_disp = raw_device_name

                            final_max_str = (
                                rec_stage_params[s_idx]["max_cap"]
                                if s_idx < len(rec_stage_params) and rec_stage_params[s_idx]["max_cap"]
                                else stg["fallback_max"]
                            )
                            final_div_str = (
                                rec_stage_params[s_idx]["div_str"]
                                if s_idx < len(rec_stage_params) and rec_stage_params[s_idx]["div_str"]
                                else stg["fallback_div"]
                            )

                            file_instruments.append(
                                {
                                    "type_kind": "힘 시험기",
                                    "cert_no": cert_no,
                                    "device_name": dev_disp,
                                    "maker": maker_name,
                                    "model": "",
                                    "serial": serial_no,
                                    "max_str": final_max_str,
                                    "div_str": final_div_str,
                                    "res_dec": get_decimal_places(stg["pts"][0]["nom"])
                                    if stg["pts"]
                                    else 0,
                                    "gubun": stg["method"],
                                    "pts": stg["pts"],
                                    "is_sub_channel": s_idx > 0,
                                }
                            )

                    # ========================================================
                    # [기기 판별 C] 온도 기록계 / 온도 지시계
                    # ========================================================
                    elif fname.startswith("T") or any(
                        k in raw_device_name for k in ["온도", "기록계", "조절계"]
                    ):
                        res_sheets = [
                            s for s in wb_src.sheetnames if "교정결과" in s and "영문" not in s
                        ]

                        def _sort_sheet_key(s):
                            m = re.search(r"\((\d+)\)", s)
                            return int(m.group(1)) if m else 1

                        res_sheets.sort(key=_sort_sheet_key)

                        if res_sheets:
                            ws_r0 = wb_src[res_sheets[0]]
                            for r_idx, row in enumerate(ws_r0.iter_rows(values_only=True)):
                                if r_idx > 18:
                                    break
                                c3 = str(row[2] if len(row) > 2 else "").replace(" ", "")
                                c11 = clean_text(row[10] if len(row) > 10 else "")
                                if "기기명" in c3 and not raw_device_name:
                                    raw_device_name = c11
                                if "제작회사" in c3 and not maker_name:
                                    parts = c11.split("/")
                                    maker_name = clean_text(parts[0])
                                    if len(parts) > 1 and not model_name:
                                        model_name = clean_text(parts[1])
                                if "기기번호" in c3 and not serial_no:
                                    serial_no = c11

                        if not raw_device_name:
                            raw_device_name = "온도 기록계(센서제외)"

                        temp_channels = []

                        for sname in res_sheets:
                            ws_temp = wb_src[sname]
                            rows_temp = list(ws_temp.iter_rows(values_only=True))

                            sheet_sensor_type = "K"
                            for r_idx in range(min(22, len(rows_temp))):
                                r_str = " ".join([str(v or "") for v in rows_temp[r_idx]])
                                if "온도교정" in r_str or "온도 교정" in r_str:
                                    m_t = re.search(r"온도\s*교정\s*[\(（](.*)", r_str)
                                    if m_t:
                                        sheet_sensor_type = clean_sensor_type(m_t.group(1))
                                        break

                            r_i = 0
                            while r_i < len(rows_temp):
                                r_vals = rows_temp[r_i]
                                r_str = " ".join([str(v or "") for v in r_vals if v is not None])

                                m_ch = re.search(r"[\(（]\s*(CH\s*\d+)\s*[\)）]", r_str, re.IGNORECASE)
                                if m_ch:
                                    ch_label = m_ch.group(1).upper().replace(" ", "")

                                    indiv_serial = ""
                                    m_sn = re.search(r"기기번호\s*:\s*([^|\(,\n]+)", r_str)
                                    if m_sn:
                                        indiv_serial = m_sn.group(1).strip()
                                        indiv_serial = clean_text(
                                            re.sub(r"\(단위.*?\)", "", indiv_serial)
                                        )

                                    ch_type = sheet_sensor_type
                                    m_local_t = re.search(r"온도\s*교정\s*[\(（](.*)", r_str)
                                    if m_local_t:
                                        ch_type = clean_sensor_type(m_local_t.group(1))

                                    nom_col = None
                                    ind_col = None
                                    data_start = None

                                    for dr in range(r_i + 1, min(r_i + 5, len(rows_temp))):
                                        d_row = rows_temp[dr]
                                        d_str = " ".join([str(v or "") for v in d_row if v is not None])
                                        if "기준온도" in d_str and "지시온도" in d_str:
                                            for c_idx, v in enumerate(d_row):
                                                if v and "기준온도" in str(v):
                                                    nom_col = c_idx
                                                if v and "지시온도" in str(v):
                                                    ind_col = c_idx
                                            data_start = dr + 1
                                            if data_start < len(rows_temp):
                                                sub_str = " ".join(
                                                    [
                                                        str(v or "")
                                                        for v in rows_temp[data_start]
                                                        if v is not None
                                                    ]
                                                )
                                                if (
                                                    "기준온도-지시온도" in sub_str
                                                    or "신뢰수준" in sub_str
                                                    or "℃" in sub_str
                                                ):
                                                    data_start += 1
                                            break

                                    pts = []
                                    if data_start and nom_col is not None and ind_col is not None:
                                        curr = data_start
                                        while curr < len(rows_temp):
                                            crow = rows_temp[curr]
                                            c_str = " ".join(
                                                [str(v or "") for v in crow if v is not None]
                                            )
                                            if any(
                                                k in c_str
                                                for k in ["끝.", "※", "권장차기교정일", "KCS-QP"]
                                            ):
                                                break
                                            if re.search(
                                                r"[\(（]\s*CH\s*\d+\s*[\)）]", c_str, re.IGNORECASE
                                            ):
                                                break
                                            if "온도교정" in c_str or "기기명" in c_str:
                                                break

                                            v_nom = crow[nom_col] if len(crow) > nom_col else None
                                            v_ind = crow[ind_col] if len(crow) > ind_col else None

                                            if v_nom is not None:
                                                try:
                                                    clean_nom = str(v_nom).replace(",", "").strip()
                                                    clean_ind = (
                                                        str(v_ind).replace(",", "").strip()
                                                        if v_ind is not None
                                                        else None
                                                    )
                                                    f_nom = float(clean_nom)
                                                    f_ind = (
                                                        float(clean_ind)
                                                        if clean_ind is not None
                                                        else None
                                                    )
                                                    pts.append(
                                                        {
                                                            "nom": f_nom,
                                                            "ind": f_ind if f_ind is not None else "-",
                                                            "ind_raw": v_ind,
                                                            "sub": "-",
                                                        }
                                                    )
                                                    if len(pts) >= 10:
                                                        break
                                                except Exception:
                                                    pass
                                            curr += 1

                                    temp_channels.append(
                                        {
                                            "ch_label": ch_label,
                                            "indiv_serial": indiv_serial,
                                            "sensor_type": ch_type,
                                            "pts": pts,
                                        }
                                    )
                                r_i += 1

                        if not temp_channels and res_sheets:
                            ws_res_single = wb_src[res_sheets[0]]
                            res_rows = list(ws_res_single.iter_rows(values_only=True))
                            fallback_type = "K"
                            for r in range(10, min(17, len(res_rows))):
                                val_t = (
                                    str(res_rows[r][1] if len(res_rows[r]) > 1 else "")
                                    + " "
                                    + str(res_rows[r][2] if len(res_rows[r]) > 2 else "")
                                )
                                if "온도교정" in val_t or "온도 교정" in val_t:
                                    m_s = re.search(r"온도\s*교정\s*[\(（](.*)", val_t)
                                    if m_s:
                                        fallback_type = clean_sensor_type(m_s.group(1))
                                    break

                            pts = []
                            for r in range(17, min(35, len(res_rows))):
                                c3_v = res_rows[r][2] if len(res_rows[r]) > 2 else None
                                c11_v = res_rows[r][10] if len(res_rows[r]) > 10 else None
                                if c3_v is not None and isinstance(c3_v, (int, float)):
                                    pts.append(
                                        {
                                            "nom": c3_v,
                                            "ind": c11_v if c11_v is not None else "-",
                                            "ind_raw": c11_v,
                                            "sub": "-",
                                        }
                                    )
                                    if len(pts) >= 10:
                                        break
                                elif str(c3_v).strip() == "끝.":
                                    break

                            temp_channels.append(
                                {
                                    "ch_label": "CH1",
                                    "indiv_serial": "",
                                    "sensor_type": fallback_type,
                                    "pts": pts,
                                }
                            )

                        total_ch = len(temp_channels)
                        for ch_idx, ch in enumerate(temp_channels):
                            res_dec = 0
                            for p in ch["pts"]:
                                raw_ind = p.get("ind_raw", None)
                                if raw_ind is not None:
                                    d = get_decimal_places(raw_ind)
                                    if d > res_dec:
                                        res_dec = d
                            div_str = f"0.{'0' * (res_dec - 1)}1 ℃" if res_dec > 0 else "1 ℃"

                            clean_base_dev = raw_device_name.replace(
                                "온도기록계", "온도 기록계"
                            ).replace("온도지시계", "온도 지시계")
                            max_channel_str = f"<{ch['ch_label']}/{total_ch}CH>" if total_ch > 1 else ""

                            final_serial = ch["indiv_serial"] if ch["indiv_serial"] else serial_no
                            if not final_serial or final_serial == "교정결과 참조":
                                final_serial = serial_no if serial_no != "교정결과 참조" else ""

                            file_instruments.append(
                                {
                                    "type_kind": "온도기록계",
                                    "cert_no": cert_no,
                                    "device_name": clean_base_dev,
                                    "maker": maker_name,
                                    "model": model_name,
                                    "serial": final_serial,
                                    "max_str": max_channel_str,
                                    "div_str": div_str,
                                    "res_dec": res_dec,
                                    "gubun": ch["sensor_type"],
                                    "pts": ch["pts"],
                                    "is_sub_channel": ch_idx > 0,
                                }
                            )

                    # ========================================================
                    # [기기 판별 D] 압력계 / 차압계
                    # ========================================================
                    elif (
                        fname.startswith("P")
                        or any(k in raw_device_name for k in ["압력", "진공", "차압"])
                        or any(u in capacity_raw for u in ["bar", "MPa", "kPa", "mmH2O"])
                    ):
                        result_sheet_names = [
                            s for s in wb_src.sheetnames if s.startswith("교정결과") and "영문" not in s
                        ]
                        ws_res = wb_src[result_sheet_names[0]] if result_sheet_names else wb_src.active

                        p_dev_name = raw_device_name
                        p_maker = maker_name
                        p_capacity = capacity_raw
                        p_serial = serial_no

                        for row in ws_res.iter_rows(values_only=True):
                            for c in row:
                                if not c:
                                    continue
                                cs = str(c).strip()
                                if ("기기명" in cs or "기   기   명" in cs) and ":" in cs and not p_dev_name:
                                    p_dev_name = clean_text(cs.split(":", 1)[1])
                                elif (
                                    ("제작회사" in cs or "제 작 회 사" in cs)
                                    and ":" in cs
                                    and not p_maker
                                ):
                                    p_maker = clean_text(cs.split(":", 1)[1])
                                elif ("형식" in cs or "형   식" in cs) and ":" in cs:
                                    val_f = clean_text(cs.split(":", 1)[1])
                                    if val_f:
                                        p_capacity = val_f
                                elif (
                                    ("기기번호" in cs or "기   기   번   호" in cs)
                                    and ":" in cs
                                    and not p_serial
                                ):
                                    p_serial = clean_text(cs.split(":", 1)[1])

                        if not p_dev_name:
                            p_dev_name = "차압계" if "차압" in capacity_raw else "게이지압용 압력계(다이얼형)"

                        p_div_str = ""
                        if "기록지" in wb_src.sheetnames:
                            ws_rec_p = wb_src["기록지"]
                            rows_rec_p = list(ws_rec_p.iter_rows(values_only=True))
                            min_scale = None
                            resolution = None
                            p_unit = ""

                            for r_i, r_vals in enumerate(rows_rec_p):
                                row_text = " ".join([str(v or "") for v in r_vals])
                                if "최소 눈금" in row_text or "최소눈금" in row_text:
                                    if r_i + 1 < len(rows_rec_p):
                                        next_r = rows_rec_p[r_i + 1]
                                        for c_i, v in enumerate(r_vals):
                                            if v and ("최소 눈금" in str(v) or "최소눈금" in str(v)):
                                                val_m = next_r[c_i]
                                                if val_m is not None and isinstance(val_m, (int, float)):
                                                    min_scale = val_m
                                                    for u_val in next_r[c_i + 1 : c_i + 12]:
                                                        if u_val and any(
                                                            u in str(u_val)
                                                            for u in [
                                                                "bar",
                                                                "MPa",
                                                                "kPa",
                                                                "Pa",
                                                                "psi",
                                                                "cmHg",
                                                                "mmHg",
                                                                "kgf",
                                                                "mmH2O",
                                                            ]
                                                        ):
                                                            p_unit = (
                                                                str(u_val)
                                                                .replace("(", "")
                                                                .replace(")", "")
                                                                .strip()
                                                            )
                                                            break
                                            if v and "분해능" in str(v):
                                                val_d = next_r[c_i]
                                                if val_d is not None and isinstance(val_d, (int, float)):
                                                    resolution = val_d
                                                    if not p_unit:
                                                        for u_val in next_r[c_i + 1 : c_i + 12]:
                                                            if u_val and any(
                                                                u in str(u_val)
                                                                for u in [
                                                                    "bar",
                                                                    "MPa",
                                                                    "kPa",
                                                                    "Pa",
                                                                    "psi",
                                                                    "cmHg",
                                                                    "mmHg",
                                                                    "kgf",
                                                                    "mmH2O",
                                                                ]
                                                            ):
                                                                p_unit = (
                                                                    str(u_val)
                                                                    .replace("(", "")
                                                                    .replace(")", "")
                                                                    .strip()
                                                                )
                                                                break
                                        if min_scale is not None and resolution is not None:
                                            break

                            if min_scale is not None and resolution is not None and resolution > 0:
                                ratio = min_scale / resolution
                                r_str = f"{int(ratio) if float(ratio).is_integer() else ratio}"
                                m_str = f"{int(min_scale) if float(min_scale).is_integer() else min_scale}"
                                p_div_str = f"{m_str}({r_str}) {p_unit}".strip()

                        pts = []
                        is_diff_pressure = ("차압" in p_dev_name) or ("차압" in raw_device_name)

                        if is_diff_pressure and "기록지" in wb_src.sheetnames:
                            ws_rec_diff = wb_src["기록지"]
                            rows_rec_all = list(ws_rec_diff.iter_rows(values_only=True))

                            col_1st = None
                            row_1st_header = None
                            press_unit = ""

                            for r_idx, r_vals in enumerate(rows_rec_all):
                                for c_idx, val in enumerate(r_vals):
                                    if val and str(val).strip() == "1회":
                                        col_1st = c_idx
                                        row_1st_header = r_idx
                                        break
                                if col_1st is not None:
                                    break

                            if col_1st is not None and row_1st_header is not None:
                                unit_row = (
                                    rows_rec_all[row_1st_header + 1]
                                    if row_1st_header + 1 < len(rows_rec_all)
                                    else []
                                )
                                col_nom = None

                                for c_back in range(col_1st - 1, -1, -1):
                                    if (
                                        c_back < len(unit_row)
                                        and unit_row[c_back] is not None
                                        and str(unit_row[c_back]).strip() != ""
                                    ):
                                        col_nom = c_back
                                        press_unit = (
                                            str(unit_row[c_back]).replace("(", "").replace(")", "").strip()
                                        )
                                        break

                                if col_nom is None:
                                    col_nom = col_1st - 1

                                if not press_unit and col_1st < len(unit_row) and unit_row[col_1st]:
                                    press_unit = (
                                        str(unit_row[col_1st]).replace("(", "").replace(")", "").strip()
                                    )

                                col_pt = 1
                                for c_check in range(0, col_nom):
                                    c_header = str(
                                        rows_rec_all[row_1st_header - 1][c_check]
                                        if row_1st_header - 1 >= 0
                                        else ""
                                    )
                                    if "측정점" in c_header:
                                        col_pt = c_check
                                        break

                                start_data_r = -1
                                for r_scan in range(
                                    row_1st_header + 1, min(row_1st_header + 6, len(rows_rec_all))
                                ):
                                    if ws_rec_diff.row_dimensions[r_scan + 1].hidden:
                                        continue
                                    v_test = (
                                        rows_rec_all[r_scan][col_1st]
                                        if len(rows_rec_all[r_scan]) > col_1st
                                        else None
                                    )
                                    try:
                                        if v_test is not None:
                                            float(str(v_test).replace(" ", "").replace(",", ""))
                                            start_data_r = r_scan
                                            break
                                    except ValueError:
                                        pass

                                raw_points = []
                                if start_data_r != -1:
                                    for r_data in range(start_data_r, len(rows_rec_all)):
                                        if ws_rec_diff.row_dimensions[r_data + 1].hidden:
                                            continue

                                        v_nom_cell = ws_rec_diff.cell(row=r_data + 1, column=col_nom + 1)
                                        v_ind_cell = ws_rec_diff.cell(row=r_data + 1, column=col_1st + 1)

                                        raw_nom = format_meaningful_decimals(v_nom_cell.value)
                                        raw_ind = get_formatted_value(v_ind_cell)

                                        v_pt = (
                                            rows_rec_all[r_data][col_pt]
                                            if len(rows_rec_all[r_data]) > col_pt
                                            else None
                                        )

                                        if raw_nom == "-" and raw_ind is None:
                                            break

                                        try:
                                            f_nom = float(
                                                str(v_nom_cell.value).replace(" ", "").replace(",", "")
                                            )
                                            f_ind = (
                                                float(
                                                    str(v_ind_cell.value).replace(" ", "").replace(",", "")
                                                )
                                                if raw_ind is not None
                                                else None
                                            )

                                            try:
                                                pt_num = (
                                                    int(float(str(v_pt).strip()))
                                                    if v_pt is not None
                                                    else len(raw_points) + 1
                                                )
                                            except Exception:
                                                pt_num = len(raw_points) + 1

                                            raw_points.append(
                                                {
                                                    "pt": pt_num,
                                                    "nom_f": f_nom,
                                                    "ind_f": f_ind,
                                                    "raw_nom": raw_nom,
                                                    "raw_ind": raw_ind if raw_ind is not None else "-",
                                                }
                                            )
                                        except Exception:
                                            break

                                if raw_points:
                                    peak_idx = 0
                                    max_val = -float("inf")
                                    for idx, item in enumerate(raw_points):
                                        val = item["nom_f"]
                                        if val >= max_val:
                                            max_val = val
                                            peak_idx = idx
                                        else:
                                            break
                                    sliced_points = raw_points[: peak_idx + 1]

                                    first_pt = sliced_points[0]
                                    if first_pt["pt"] == 1 and abs(first_pt["nom_f"]) < 1e-7:
                                        sliced_points = sliced_points[1:]

                                    for p_i, item in enumerate(sliced_points[:10]):
                                        nom_final = (
                                            f"{item['raw_nom']} {press_unit}"
                                            if p_i == 0 and press_unit
                                            else item["raw_nom"]
                                        )
                                        ind_final = (
                                            f"{item['raw_ind']} {press_unit}"
                                            if p_i == 0 and press_unit and item["raw_ind"] != "-"
                                            else item["raw_ind"]
                                        )

                                        pts.append({"nom": nom_final, "ind": ind_final, "sub": "-"})

                        else:
                            rows_res_p = list(ws_res.iter_rows(values_only=True))
                            data_start_r = -1
                            press_unit = "MPa"
                            std_col = 15
                            ind_col = 35

                            for r_i, r_vals in enumerate(rows_res_p):
                                row_str = " ".join([str(v or "") for v in r_vals])
                                if ("표준" in row_str and "압력" in row_str and "지시" in row_str) or (
                                    "표준기" in row_str and "지시" in row_str
                                ):
                                    for c_i, v in enumerate(r_vals):
                                        if v and "표준" in str(v) and "압력" in str(v):
                                            std_col = c_i
                                        elif v and "지시" in str(v):
                                            ind_col = c_i

                                    for offset in [1, 2, 3]:
                                        if r_i + offset < len(rows_res_p):
                                            unit_row_str = " ".join(
                                                [str(v or "") for v in rows_res_p[r_i + offset]]
                                            )
                                            for u in [
                                                "MPa",
                                                "kPa",
                                                "bar",
                                                "psi",
                                                "cmHg",
                                                "mmHg",
                                                "Pa",
                                                "kgf",
                                            ]:
                                                if u in unit_row_str:
                                                    press_unit = u
                                                    break
                                            if press_unit != "MPa":
                                                break

                                    for check_r in range(r_i + 1, min(r_i + 6, len(rows_res_p))):
                                        if ws_res.row_dimensions[check_r + 1].hidden:
                                            continue
                                        v_check = (
                                            rows_res_p[check_r][std_col]
                                            if len(rows_res_p[check_r]) > std_col
                                            else None
                                        )
                                        try:
                                            if v_check is not None:
                                                float(str(v_check).replace(",", ""))
                                                data_start_r = check_r
                                                break
                                        except ValueError:
                                            pass
                                    break

                            if data_start_r != -1:
                                for r_i in range(data_start_r, min(data_start_r + 25, len(rows_res_p))):
                                    if ws_res.row_dimensions[r_i + 1].hidden:
                                        continue

                                    crow = rows_res_p[r_i]
                                    crow_str = " ".join([str(v or "") for v in crow])
                                    if any(k in crow_str for k in ["*", "※", "끝."]):
                                        break

                                    std_cell = ws_res.cell(row=r_i + 1, column=std_col + 1)
                                    ind_cell = ws_res.cell(row=r_i + 1, column=ind_col + 1)

                                    if std_cell.value is not None:
                                        try:
                                            f_std = float(str(std_cell.value).replace(",", ""))
                                            if abs(f_std) < 1e-7 and len(pts) == 0:
                                                continue

                                            std_s = format_meaningful_decimals(std_cell.value)
                                            if std_s != "-":
                                                std_s = f"{std_s} {press_unit}".strip()

                                            ind_s = get_formatted_value(ind_cell)
                                            if ind_s is not None:
                                                ind_s = f"{ind_s} {press_unit}".strip()
                                            else:
                                                ind_s = "-"

                                            pts.append({"nom": std_s, "ind": ind_s, "sub": "-"})
                                            if len(pts) >= 10:
                                                break
                                        except Exception:
                                            pass

                        file_instruments.append(
                            {
                                "type_kind": "압력계",
                                "cert_no": cert_no,
                                "device_name": p_dev_name.replace("게이지압용 ", "게이지압용\n"),
                                "maker": p_maker,
                                "model": "",
                                "serial": p_serial,
                                "max_str": p_capacity,
                                "div_str": p_div_str,
                                "res_dec": 0,
                                "gubun": "",
                                "pts": pts,
                            }
                        )

                    # ========================================================
                    # [기기 판별 E] 전기식 지시 저울 (기록지 탭 -> 2.2 직선성 시험 기준)
                    # ========================================================
                    else:
                        if "교정결과" in wb_src.sheetnames:
                            ws_res = wb_src["교정결과"]
                            rows_res = list(ws_res.iter_rows(values_only=True))
                            for r in range(0, min(20, len(rows_res))):
                                c4 = str(rows_res[r][3] if len(rows_res[r]) > 3 else "").replace(" ", "")
                                c10 = clean_text(rows_res[r][9] if len(rows_res[r]) > 9 else "")
                                if "제작회사" in c4 and not maker_name:
                                    maker_name = clean_text(c10.split("/")[0])
                                if "기기번호" in c4 and not serial_no:
                                    serial_no = c10
                                if "최대용량" in c4 and not capacity_raw:
                                    capacity_raw = c10
                                if "분해능" in c4 and not scale_raw:
                                    scale_raw = c10

                        gubun_val = "크레인" if "크레인" in raw_device_name else ""
                        max_str, div_str, _ = format_capacity_and_resolution_robust(
                            capacity_raw, scale_raw
                        )

                        pts = []

                        # 1순위: '기록지' 탭의 '2.2 직선성시험'에서 명목값(범위) 및 평균값(측정결과) 추출
                        if "기록지" in wb_src.sheetnames:
                            ws_rec = wb_src["기록지"]
                            nom_col = 2
                            avg_col = 7
                            data_start_row = None

                            for r in range(1, ws_rec.max_row + 1):
                                row_vals = [ws_rec.cell(r, c).value for c in range(1, 15)]
                                row_str = " ".join([str(v) for v in row_vals if v is not None])
                                if "직선성" in row_str:
                                    for hr in range(r + 1, min(r + 6, ws_rec.max_row + 1)):
                                        for hc in range(1, 15):
                                            hv = str(ws_rec.cell(hr, hc).value or "").strip()
                                            if "명목값" in hv:
                                                nom_col = hc
                                            elif "평균값" in hv or hv == "평균":
                                                avg_col = hc
                                                data_start_row = hr + 1
                                    if data_start_row is None:
                                        data_start_row = r + 3
                                    break

                            if data_start_row is not None:
                                first_valid_row_checked = False
                                for r in range(data_start_row, ws_rec.max_row + 1):
                                    if ws_rec.row_dimensions[r].hidden:
                                        continue

                                    row_vals = [ws_rec.cell(r, c).value for c in range(1, 12)]
                                    row_str = " ".join([str(v) for v in row_vals if v is not None])
                                    if any(
                                        k in row_str
                                        for k in ["3.불확도", "3. 불확도", "불확도 계산", "*UA", "끝."]
                                    ):
                                        break

                                    nom_cell = ws_rec.cell(r, nom_col)
                                    avg_cell = ws_rec.cell(r, avg_col)

                                    # 명목값이 비어 있는 행(예: 수식만 남은 마지막 빈 행)은 건너뜀
                                    if nom_cell.value is None or str(nom_cell.value).strip() == "":
                                        continue

                                    try:
                                        f_nom = float(str(nom_cell.value).replace(",", "").strip())
                                    except ValueError:
                                        continue

                                    # 직선성 시험에서 제일 처음 최초 시작되는 명목값이 0일 때만 제외
                                    if not first_valid_row_checked:
                                        first_valid_row_checked = True
                                        if abs(f_nom) < 1e-9:
                                            continue

                                    nom_str = get_formatted_value(nom_cell)
                                    avg_str = get_formatted_value(avg_cell)

                                    pts.append(
                                        {
                                            "nom": nom_str if nom_str is not None else "-",
                                            "ind": avg_str if avg_str is not None else "-",
                                            "sub": "-",
                                        }
                                    )
                                    if len(pts) >= 10:
                                        break

                        # 2순위 폴백: 혹시라도 '기록지' 탭이 없는 특이 양식일 경우 '교정결과' 탭에서 서식 유지 추출
                        if not pts and "교정결과" in wb_src.sheetnames:
                            ws_res = wb_src["교정결과"]
                            start_r = 0
                            for r in range(1, min(26, ws_res.max_row + 1)):
                                if "저울지시값" in str(ws_res.cell(r, 3).value or ""):
                                    start_r = r + 1
                                    break

                            if start_r > 0:
                                first_checked = False
                                for r in range(start_r, min(start_r + 20, ws_res.max_row + 1)):
                                    ind_cell = ws_res.cell(r, 3)
                                    std_cell = ws_res.cell(r, 9)
                                    val_str = str(ind_cell.value or "")
                                    if "※" in val_str or "끝" in val_str:
                                        break
                                    if ind_cell.value is None and std_cell.value is None:
                                        continue
                                    try:
                                        f_ind = float(str(ind_cell.value).replace(",", "").strip())
                                    except ValueError:
                                        continue
                                    if not first_checked:
                                        first_checked = True
                                        if abs(f_ind) < 1e-9:
                                            continue
                                    pts.append(
                                        {
                                            "nom": format_meaningful_decimals(std_cell.value)
                                            if std_cell.value is not None
                                            else format_meaningful_decimals(ind_cell.value),
                                            "ind": get_formatted_value(ind_cell) or "-",
                                            "sub": "-",
                                        }
                                    )
                                    if len(pts) >= 10:
                                        break

                        dev_disp = (
                            "전기식 지시 저울\n(크레인)"
                            if "크레인" in raw_device_name
                            else raw_device_name.replace("전기식지시저울", "전기식 지시 저울")
                        )
                        if not dev_disp:
                            dev_disp = "전기식 지시 저울"

                        file_instruments.append(
                            {
                                "type_kind": "전기식 저울",
                                "cert_no": cert_no,
                                "device_name": dev_disp,
                                "maker": maker_name,
                                "model": "",
                                "serial": serial_no,
                                "max_str": max_str,
                                "div_str": div_str,
                                "res_dec": 0,
                                "gubun": gubun_val,
                                "pts": pts,
                            }
                        )

                    for sub_i, inst in enumerate(file_instruments):
                        inst["cert_file_no"] = file_idx if sub_i == 0 else ""
                        all_instruments.append(inst)

                except Exception as file_err:
                    error_files.append(f"{file.name} ({file_err})")
                finally:
                    if wb_src is not None:
                        try:
                            wb_src.close()
                        except Exception:
                            pass

            if error_files:
                st.warning(
                    "⚠️ 일부 파일을 처리하는 중 오류가 발생했습니다:\n- " + "\n- ".join(error_files)
                )

            if all_instruments:
                # 1. 각 기기별 페이지 및 슬롯 위치 사전 배정 (2칸짜리 기기가 페이지 경계에 걸리면 다음 페이지로 이월)
                placements = []
                curr_slot = 0
                total_slots_used = 0

                for inst in all_instruments:
                    slots_for_inst = 2 if len(inst["pts"]) > 5 else 1
                    total_slots_used += slots_for_inst

                    if (curr_slot % MAX_SLOTS_PER_PAGE) + slots_for_inst > MAX_SLOTS_PER_PAGE:
                        curr_slot = ((curr_slot // MAX_SLOTS_PER_PAGE) + 1) * MAX_SLOTS_PER_PAGE

                    page_idx = curr_slot // MAX_SLOTS_PER_PAGE
                    slot_in_page = curr_slot % MAX_SLOTS_PER_PAGE
                    placements.append((inst, page_idx, slot_in_page))
                    curr_slot += slots_for_inst

                total_pages_needed = max(1, (curr_slot + MAX_SLOTS_PER_PAGE - 1) // MAX_SLOTS_PER_PAGE)
                max_row_limit = 3 + MAX_SLOTS_PER_PAGE * 6

                wb_out = openpyxl.load_workbook(TEMPLATE_FILE)
                ws_template = wb_out["교정점검표"]

                sheet_list = [ws_template]
                for p in range(2, total_pages_needed + 1):
                    ws_new = wb_out.copy_worksheet(ws_template)
                    ws_new.title = f"교정점검표 ({p})"
                    sheet_list.append(ws_new)

                for ws_target in sheet_list:
                    if global_client_name:
                        ws_target["C1"] = str(global_client_name)
                        ws_target["C1"].alignment = Alignment(horizontal="left", vertical="center")

                    ws_target["H1"] = "-"
                    ws_target["H1"].alignment = Alignment(horizontal="center", vertical="center")
                    ws_target["H2"] = "-"
                    ws_target["H2"].alignment = Alignment(horizontal="center", vertical="center")
                    ws_target["I2"] = "   月    日"
                    ws_target["I2"].alignment = Alignment(horizontal="center", vertical="center")

                    for slot_i in range(MAX_SLOTS_PER_PAGE):
                        ws_target.cell(row=4 + slot_i * 6, column=1, value="")

                preview_rows = []

                for inst, page_idx, slot_in_page in placements:
                    pts = inst["pts"]
                    res_dec = inst["res_dec"]
                    ws_curr = sheet_list[page_idx]

                    # 신규 양식: 1슬롯당 6행 차지 (4행, 10행, 16행 ...)
                    start_row = 4 + slot_in_page * 6

                    # A열: No
                    if not del_number and inst["cert_file_no"]:
                        ws_curr.cell(row=start_row, column=1, value=inst["cert_file_no"])

                    # B열: 기기명 (서브채널은 기기명만 비움)
                    dev_name_to_write = (
                        "" if inst.get("is_sub_channel") else str(inst["device_name"])
                    )
                    cell_dev = ws_curr.cell(row=start_row, column=2, value=dev_name_to_write)
                    cell_dev.alignment = Alignment(
                        horizontal="center", vertical="center", wrap_text=True
                    )

                    # C열: 제조사(상단 3행 병합) / 모델(하단 3행 병합 -> start_row + 3)
                    maker_to_write = str(inst["maker"])
                    if inst.get("is_sub_channel"):
                        model_to_write = ""
                    else:
                        model_to_write = (
                            ""
                            if inst["type_kind"] in ["압력계", "힘 시험기"]
                            else str(inst["model"])
                        )

                    ws_curr.cell(row=start_row, column=3, value=maker_to_write)
                    ws_curr.cell(row=start_row + 3, column=3, value=model_to_write)

                    # D열: 기기번호
                    ws_curr.cell(row=start_row, column=4, value=str(inst["serial"]))

                    # E열: MAX(상단 3행 병합) / DIV(하단 3행 병합 -> start_row + 3)
                    ws_curr.cell(row=start_row, column=5, value=str(inst["max_str"]))
                    ws_curr.cell(row=start_row + 3, column=5, value=str(inst["div_str"]))

                    # F열: 방식 / 구분 / 타입
                    ws_curr.cell(row=start_row, column=6, value=str(inst["gubun"]))

                    # G열 & H열: 범위 및 측정결과 (신규 6행 슬롯 내 병합 구조 대응)
                    nom_preview_list = []
                    ind_preview_list = []

                    for p_idx, pt in enumerate(pts[:10]):
                        target_r = start_row + PT_ROW_OFFSETS[p_idx]
                        if target_r > max_row_limit:
                            break

                        if isinstance(pt["nom"], str):
                            nom_val_str = pt["nom"]
                        else:
                            nom_val_str = format_number(pt["nom"], res_dec)

                        if isinstance(pt["ind"], str):
                            ind_val_str = pt["ind"]
                        else:
                            ind_base = format_number(pt["ind"], res_dec)
                            sub_val = pt.get("sub", "-")
                            if sub_val != "-" and sub_val != 0 and sub_val != "0":
                                ind_val_str = f"{ind_base} (가:{format_sub_weight(sub_val)})"
                            else:
                                ind_val_str = ind_base

                        nom_preview_list.append(nom_val_str)
                        ind_preview_list.append(ind_val_str)

                        if not del_nominal:
                            ws_curr.cell(row=target_r, column=7, value=nom_val_str)
                        if not del_data:
                            ws_curr.cell(row=target_r, column=8, value=ind_val_str)

                    # I열: 교정번호 공란
                    ws_curr.cell(row=start_row, column=9, value="")

                    preview_rows.append(
                        {
                            "No": inst["cert_file_no"] if inst["cert_file_no"] else "-",
                            "분류": inst["type_kind"],
                            "기기명": dev_name_to_write.replace("\n", " ")
                            if dev_name_to_write
                            else "(상동/서브채널)",
                            "제조사": maker_to_write,
                            "기기번호": inst["serial"],
                            "MAX / DIV": f"{inst['max_str']} / {inst['div_str']}".strip(" /"),
                            "※(구분)": inst["gubun"] if inst["gubun"] else "-",
                            "범위 (명목값)": " , ".join(nom_preview_list),
                            "측정결과 (지시/평균값)": " , ".join(ind_preview_list),
                        }
                    )

                output_buffer = io.BytesIO()
                wb_out.save(output_buffer)
                output_buffer.seek(0)
                wb_out.close()

                safe_client = re.sub(r'[\/:*?"<>|() ]', "", str(global_client_name)) or "미상업체"
                out_name = f"기록지({MAX_SLOTS_PER_PAGE}칸)_{safe_client}_총{len(uploaded_files)}건_({total_pages_needed}P).xlsx"
                
                # ========================================================
                # 다운로드 버튼 렌더링 (좌측 컨트롤 패널 영역으로 이동)
                # ========================================================
                with download_placeholder:
                    st.markdown("<div style='margin-top: 0.85rem;'></div>", unsafe_allow_html=True)
                    st.download_button(
                        label=f"📥 [딸깍] {out_name} 다운로드",
                        data=output_buffer,
                        file_name=out_name,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="primary",
                        use_container_width=True,
                    )

                # ========================================================
                # [결과 대시보드 및 실시간 미리보기 렌더링]
                # ========================================================
                with bottom_result_container:
                    st.markdown("<div style='margin-top: 1.5rem;'></div>", unsafe_allow_html=True)
                    with st.container(border=True):
                        st.markdown(
                            f"""
                            <div class="section-header">
                                <div class="section-title">✅ {MAX_SLOTS_PER_PAGE}칸 기록지 변환 완료</div>
                            </div>
                            """,
                            unsafe_allow_html=True,
                        )

                        m_c1, m_c2, m_c3, m_c4 = st.columns(4, gap="small")
                        m_c1.metric(
                            "🏢 의뢰 업체명",
                            str(global_client_name) if global_client_name else "미확인",
                        )
                        m_c2.metric("📑 처리된 성적서", f"{len(uploaded_files)} 건")
                        m_c3.metric("🎯 사용 슬롯", f"{total_slots_used} 칸 ({MAX_SLOTS_PER_PAGE}칸 양식)")
                        m_c4.metric("📄 생성 페이지", f"{total_pages_needed} 페이지")

                        st.markdown("<div style='margin-top: 1rem;'></div>", unsafe_allow_html=True)
                        
                        with st.expander("🔍 추출 데이터 실시간 미리보기 (클릭하여 검수)", expanded=True):
                            df_preview = pd.DataFrame(preview_rows)
                            st.dataframe(
                                df_preview,
                                use_container_width=True,
                                hide_index=True,
                            )

        except Exception as e:
            st.error(f"❌ 분석 중 예기치 않은 오류가 발생했습니다: {e}")

# ----------------- 하단 푸터 -----------------
st.markdown("<div style='margin-top: 3rem;'></div>", unsafe_allow_html=True)
st.caption("⚡ **KCS 딸깍 기록지** | Crafted by **D** | Automated Calibration Record Generator")
